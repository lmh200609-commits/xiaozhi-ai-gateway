import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
import docx
import httpx
from pathlib import Path

# 1. Create a real sample Word document
doc = docx.Document()
doc.add_heading("小智AI开发板深度使用手册与常见故障排除", 0)

doc.add_heading("一、网络与热点连接", level=1)
doc.add_paragraph(
    "小智设备支持2.4GHz Wi-Fi网络。默认出厂预置的测试热点名称为vivo X200s，"
    "热点密码为lmh200609。若需要更换WiFi，可在开机时按住BOOT按键5秒进入AP配网模式，"
    "手机连接小智广播的热点后访问192.168.4.1即可输入新WiFi。"
)

doc.add_heading("二、语音唤醒与灵敏度", level=1)
doc.add_paragraph(
    "唤醒词为“你好小智”。如果在嘈杂环境下唤醒失败，可以直接长按开发板右侧的实体BOOT按键开启语音录制，"
    "松开按键后小智会立即开始思考并回答。建议说话时距离板载麦克风10到50厘米之间效果最佳。"
)

doc.add_heading("三、声音杂音与爆音排查", level=1)
doc.add_paragraph(
    "如果听到板载扬声器播放时有轻微破音或卡顿，通常是供电电流不足所致。"
    "小智ESP32-S3在WiFi全功率射频发射与扬声器工作时瞬间峰值电流可达1.5A以上，"
    "请务必使用5V/2A以上的独立Type-C充电器供电，避免插在电脑低功率USB扩展坞上。"
)

docx_path = Path(__file__).parent / "test_manual.docx"
doc.save(docx_path)
print(f"[*] Sample docx created: {docx_path}")

# 2. Upload to our gateway API
print("[*] Uploading to http://127.0.0.1:8001/api/knowledge/upload ...")
with open(docx_path, "rb") as f:
    files = {"file": ("test_manual.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp = httpx.post("http://127.0.0.1:8001/api/knowledge/upload", files=files, timeout=60.0)
    print(f"[*] Response status: {resp.status_code}")
    data = resp.json()
    print("=== AI Structuring Result ===")
    print(f"Title: {data.get('title')}")
    print(f"Category: {data.get('category')}")
    print(f"Summary: {data.get('summary')}")
    print(f"QA Pairs Count: {data.get('qa_count')}")
    print(f"Fact Chunks Count: {data.get('fact_count')}")

# 3. Test searching against the newly added knowledge
print("\n[*] Testing search for: '喇叭声音卡顿破音怎么办？'")
search_resp = httpx.post("http://127.0.0.1:8001/api/knowledge/search", json={"query": "喇叭声音卡顿破音怎么办？"}, timeout=10.0)
sdata = search_resp.json()
print(f"[*] Search cost: {sdata.get('cost_ms')}ms, Matches: {len(sdata.get('results', []))}")
for idx, item in enumerate(sdata.get("results", [])):
    print(f"  [{idx+1}] Score: {item['score']} | {item.get('question') or item.get('title')}")
    print(f"      Content: {item['content']}")
